"""Add two comparison sheets to eval/outputs/metrics_summary.xlsx:

  "Full vs Staged" — production full-study pipeline vs single-call prompt arm, SCALAR forms
  "Full vs Table"  — production full-study pipeline vs single-call prompt arm, TABLE forms

Reads the already-scored per-field `all_forms_metrics.xlsx` under
outputs/{full_studies,staged,table}/<dataset>/<model>/, aggregates each form to macro/micro F1
(identically on both sides), and writes the two sheets WITHOUT clobbering the existing sheets
(backs up first). Pure read+aggregate+write — no re-scoring, no LLM.

    python -m eval.engine.reports.build_full_vs_prompt
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

sys.path.insert(0, "/home/ubuntu/evistream/eval")

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

EVAL_ROOT = Path("/home/ubuntu/evistream/eval")
OUT = EVAL_ROOT / "outputs"
SUMMARY = OUT / "metrics_summary.xlsx"

DATASETS = ["oral_cancer", "antibiotic", "periodontitis", "ibuprofen"]
MODELS = ["claude", "gpt", "gemini"]
LABELS = {
    "study_characteristics": "Study Characteristics", "patient_population": "Patient Population",
    "reference_standard": "Reference Standard", "index_test": "Index Test",
    "abx_study_char": "Study Characteristics", "abx_patient_pop": "Patient Population",
    "abx_interventions": "Interventions", "abx_outcomes": "Dichotomous Outcomes",
    "perio_study_char": "Study Characteristics", "perio_patient_pop": "Patient Population",
    "perio_interventions": "Interventions", "perio_outcomes": "Continuous Outcomes",
    "ibu_study_char": "Study Characteristics", "ibu_patient_pop": "Patient Population",
    "ibu_interventions": "Interventions", "ibu_continuous_outcomes": "Continuous Outcomes",
    "ibu_dichotomous_outcomes": "Dichotomous Outcomes", "ibu_risk_of_bias": "Risk of Bias",
}
DS_LABEL = {"oral_cancer": "Oral Cancer", "antibiotic": "Antibiotic", "periodontitis": "Periodontitis",
            "ibuprofen": "Ibuprofen"}
# "Summary" is a per-form rollup sheet metrics_writer.py keeps first in the workbook —
# not itself a form; sheet-name enumeration below must skip it everywhere.
NON_FORM_SHEETS = {"Summary"}

HEADERS = ["Dataset", "Form", "Model", "Fields (full)", "Fields (prompt)",
           "F1 (full)", "F1 (prompt)", "ΔF1 (full−prompt)",
           "microF1 (full)", "microF1 (prompt)", "Prec (full)", "Prec (prompt)",
           "Recall (full)", "Recall (prompt)", "TP/FP/FN (full)", "TP/FP/FN (prompt)"]
_COL = {h: i + 1 for i, h in enumerate(HEADERS)}   # 1-based column index by header

_SHEETS_CACHE: dict = {}


def _sheets(path: Path):
    if path not in _SHEETS_CACHE:
        names = pd.ExcelFile(path).sheet_names if path.exists() else []
        _SHEETS_CACHE[path] = [s for s in names if s not in NON_FORM_SHEETS]
    return _SHEETS_CACHE[path]


def _agg(path: Path, sheet: str) -> dict:
    """Form-level aggregate from a per-field metrics sheet (same method both sides)."""
    df = pd.read_excel(path, sheet_name=sheet)
    f1 = pd.to_numeric(df["F1"], errors="coerce")
    pr = pd.to_numeric(df["Precision"], errors="coerce")
    rc = pd.to_numeric(df["Recall"], errors="coerce")
    tp = pd.to_numeric(df["TP"], errors="coerce").fillna(0).sum()
    fp = pd.to_numeric(df["FP"], errors="coerce").fillna(0).sum()
    fn = pd.to_numeric(df["FN"], errors="coerce").fillna(0).sum()
    mp = tp / (tp + fp) if (tp + fp) else 0.0
    mr = tp / (tp + fn) if (tp + fn) else 0.0
    micro = 2 * mp * mr / (mp + mr) if (mp + mr) else 0.0
    return {"fields": int(len(df)), "f1": float(f1.mean()), "prec": float(pr.mean()),
            "recall": float(rc.mean()), "micro_f1": float(micro),
            "tp": int(tp), "fp": int(fp), "fn": int(fn)}


def _collect(variant: str):
    rows, gaps = [], []
    for ds in DATASETS:
        sheets = next((_sheets(OUT / variant / ds / m / "all_forms_metrics.xlsx")
                       for m in MODELS if _sheets(OUT / variant / ds / m / "all_forms_metrics.xlsx")), [])
        for sheet in sheets:
            for m in MODELS:
                pp = OUT / variant / ds / m / "all_forms_metrics.xlsx"
                fp = OUT / "full_studies" / ds / m / "all_forms_metrics.xlsx"
                if sheet not in _sheets(pp):
                    continue
                if sheet not in _sheets(fp):
                    gaps.append(f"{ds}/{m}/{sheet}")
                    continue
                rows.append({"ds": ds, "sheet": sheet, "model": m,
                             "full": _agg(fp, sheet), "prompt": _agg(pp, sheet)})
    return rows, gaps


# styling
TITLE_FT = Font(bold=True, size=12)
HDR_FT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="4472C4")
OVR_FILL = PatternFill("solid", fgColor="DDEBF7")
F_GOOD = PatternFill("solid", fgColor="C6EFCE")   # prompt ~ as good / better
F_MILD = PatternFill("solid", fgColor="FFEB9C")
F_BAD = PatternFill("solid", fgColor="F8CBAD")
CTR = Alignment(horizontal="center")


def _delta_fill(d):
    if d <= 0.02:
        return F_GOOD
    if d <= 0.07:
        return F_MILD
    return F_BAD


def _row_values(ds, form, model, full, prompt):
    d = round(full["f1"] - prompt["f1"], 3)
    return [DS_LABEL.get(ds, ds), form, model, full["fields"], prompt["fields"],
            round(full["f1"], 3), round(prompt["f1"], 3), d,
            round(full["micro_f1"], 3), round(prompt["micro_f1"], 3),
            round(full["prec"], 3), round(prompt["prec"], 3),
            round(full["recall"], 3), round(prompt["recall"], 3),
            f"{full['tp']}/{full['fp']}/{full['fn']}", f"{prompt['tp']}/{prompt['fp']}/{prompt['fn']}"], d


def _write_sheet(wb, name, variant_label, rows, gaps):
    if name in wb.sheetnames:
        wb.remove(wb[name])
    ws = wb.create_sheet(name)
    ws.cell(1, 1, f"Full study (production) vs {variant_label} prompt arm (single-call) — macro/micro F1 per form")
    ws.cell(1, 1).font = TITLE_FT
    ws.cell(3, 1, "SCORECARD — F1 / microF1 / Precision / Recall per (dataset, form, model);  ΔF1 = full − prompt  (＋ = production better)")
    ws.cell(3, 1).font = Font(italic=True)
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(4, c, h); cell.font = HDR_FT; cell.fill = HDR_FILL; cell.alignment = CTR

    r = 5
    all_deltas = []
    by_ds = {}
    for row in rows:
        by_ds.setdefault(row["ds"], []).append(row)
    for ds in DATASETS:
        ds_rows = by_ds.get(ds, [])
        if not ds_rows:
            continue
        ds_deltas = []
        for row in ds_rows:
            vals, d = _row_values(ds, LABELS.get(row["sheet"], row["sheet"]), row["model"],
                                  row["full"], row["prompt"])
            for c, v in enumerate(vals, 1):
                ws.cell(r, c, v)
            dc = _COL["ΔF1 (full−prompt)"]
            ws.cell(r, dc).fill = _delta_fill(d); ws.cell(r, dc).alignment = CTR
            ds_deltas.append(d); all_deltas.append(d)
            r += 1
        # per-dataset OVERALL (mean Δ + mean F1s)
        mf = round(sum(x["full"]["f1"] for x in ds_rows) / len(ds_rows), 3)
        mp = round(sum(x["prompt"]["f1"] for x in ds_rows) / len(ds_rows), 3)
        ws.cell(r, 1, f"{DS_LABEL.get(ds, ds)} — OVERALL (mean)").font = Font(bold=True)
        ws.cell(r, _COL["F1 (full)"], mf); ws.cell(r, _COL["F1 (prompt)"], mp)
        ws.cell(r, _COL["ΔF1 (full−prompt)"], round(mf - mp, 3))
        for c in range(1, len(HEADERS) + 1):
            ws.cell(r, c).fill = OVR_FILL
        r += 1

    # footnote
    r += 1
    note = ("Note: full_studies has all 3 models for oral_cancer but only claude for antibiotic & "
            "periodontitis — rows shown only where both sides exist.")
    if gaps:
        note += f"  No full_studies counterpart for: {', '.join(sorted(set(gaps)))}."
    ws.cell(r, 1, note).font = Font(italic=True, size=9)

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 8
    for c in range(4, len(HEADERS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 13
    ws.freeze_panes = "D5"
    return len(rows)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--summary", default=str(SUMMARY))
    args = ap.parse_args()
    summary = Path(args.summary)
    if not summary.exists():
        sys.exit(f"ERROR: {summary} not found")

    bak = summary.with_name("metrics_summary.bak.xlsx")
    shutil.copy2(summary, bak)
    print(f"  backed up → {bak}")

    staged_rows, staged_gaps = _collect("staged")
    table_rows, table_gaps = _collect("table")

    wb = openpyxl.load_workbook(summary)
    before = list(wb.sheetnames)
    n1 = _write_sheet(wb, "Full vs Staged", "Staged", staged_rows, staged_gaps)
    n2 = _write_sheet(wb, "Full vs Table", "Table", table_rows, table_gaps)
    wb.save(summary)

    print(f"  existing sheets preserved: {before}")
    print(f"  + 'Full vs Staged' ({n1} rows), 'Full vs Table' ({n2} rows)")
    print(f"  saved → {summary}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
