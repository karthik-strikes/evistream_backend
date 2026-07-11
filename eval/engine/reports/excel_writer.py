"""
Write four-sheet Excel reports per form.
Sheet 1: Side-by-side comparison (color-coded)
Sheet 2: Agreement statistics per field
Sheet 3: Discrepancy report (mismatches + potential over-extractions)
Sheet 4: Confusion matrices per field
"""

import os
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

GREEN  = PatternFill("solid", fgColor="C6EFCE")
RED    = PatternFill("solid", fgColor="FFC7CE")
ORANGE = PatternFill("solid", fgColor="FFD966")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
GREY   = PatternFill("solid", fgColor="D9D9D9")
HEADER = PatternFill("solid", fgColor="1F4E79")
HDR_FT = Font(color="FFFFFF", bold=True)
WRAP   = Alignment(wrap_text=True, vertical="top")

_DEFAULT_OUTPUTS_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "outputs", "per_form_reports")


def _outputs_dir() -> str:
    return os.environ.get("EVAL_REPORTS_DIR", _DEFAULT_OUTPUTS_DIR)


def _style_header(ws, ncols: int):
    for col_idx in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER
        cell.font = HDR_FT
        cell.alignment = WRAP


def _autowidth(ws, df: pd.DataFrame, max_width: int = 50):
    for col_idx, col_name in enumerate(df.columns, 1):
        col_vals = df[col_name].astype(str)
        max_len  = max(len(str(col_name)), col_vals.str.len().max())
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, max_width)


def write_form_report(
    form_name: str,
    comparison_df: pd.DataFrame,
    stats_df: pd.DataFrame,
    discrepancy_df: pd.DataFrame,
    confusion_data_by_field: dict = None,
):
    outputs_dir = _outputs_dir()
    os.makedirs(outputs_dir, exist_ok=True)
    out_path = os.path.join(outputs_dir, f"{form_name}_evaluation.xlsx")

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # ── Sheet 1: Comparison ────────────────────────────────────────────────
        comparison_df.to_excel(writer, index=False, sheet_name="Comparison")
        ws1 = writer.sheets["Comparison"]
        _style_header(ws1, len(comparison_df.columns))
        _autowidth(ws1, comparison_df)
        ws1.freeze_panes = "A2"

        # Color MATCH_* columns
        match_cols = [i + 1 for i, c in enumerate(comparison_df.columns)
                      if c.startswith("MATCH_")]
        for row_idx in range(2, len(comparison_df) + 2):
            for col_idx in match_cols:
                cell = ws1.cell(row=row_idx, column=col_idx)
                v = str(cell.value or "")
                if v == "YES":   cell.fill = GREEN
                elif v == "NO":  cell.fill = RED
                elif v == "NA":  cell.fill = GREY
                elif v == "SKIP": cell.fill = GREY

        # ── Sheet 2: Agreement Stats ───────────────────────────────────────────
        stats_df.to_excel(writer, index=False, sheet_name="Agreement Stats")
        ws2 = writer.sheets["Agreement Stats"]
        _style_header(ws2, len(stats_df.columns))
        _autowidth(ws2, stats_df, max_width=30)
        ws2.freeze_panes = "A2"

        # Color % agreement
        if "pct_agreement" in stats_df.columns:
            pct_col = list(stats_df.columns).index("pct_agreement") + 1
            for row_idx in range(2, len(stats_df) + 2):
                cell = ws2.cell(row=row_idx, column=pct_col)
                try:
                    v = float(cell.value)
                    cell.fill = GREEN if v >= 80 else (YELLOW if v >= 60 else RED)
                except (TypeError, ValueError):
                    pass

        # ── Sheet 3: Discrepancies ─────────────────────────────────────────────
        discrepancy_df.to_excel(writer, index=False, sheet_name="Discrepancies")
        ws3 = writer.sheets["Discrepancies"]
        _style_header(ws3, len(discrepancy_df.columns))
        _autowidth(ws3, discrepancy_df)
        ws3.freeze_panes = "A2"

        # Color rows by type: mismatch=red, potential_over_extraction=orange
        if "type" in discrepancy_df.columns:
            type_col_idx = list(discrepancy_df.columns).index("type") + 1
            n_disc_cols  = len(discrepancy_df.columns)
            for row_idx in range(2, len(discrepancy_df) + 2):
                row_type = str(ws3.cell(row=row_idx, column=type_col_idx).value or "")
                fill = RED if row_type == "mismatch" else (ORANGE if row_type == "potential_over_extraction" else None)
                if fill:
                    for col_idx in range(1, n_disc_cols + 1):
                        ws3.cell(row=row_idx, column=col_idx).fill = fill

        # ── Sheet 4: Confusion Matrices ────────────────────────────────────────
        if confusion_data_by_field:
            ws4 = writer.book.create_sheet("Confusion Matrices")
            current_row = 1
            for field_name, cd in confusion_data_by_field.items():
                labels = cd["labels"]
                y_true = cd["y_true"]
                y_pred = cd["y_pred"]

                # Field header
                header_cell = ws4.cell(row=current_row, column=1, value=f"Field: {field_name}")
                header_cell.fill = HEADER
                header_cell.font = HDR_FT
                current_row += 1

                # Column headers (AI predicted values)
                ws4.cell(row=current_row, column=1, value="GT \\ AI").fill = GREY
                for j, lbl in enumerate(labels, 2):
                    c = ws4.cell(row=current_row, column=j, value=lbl)
                    c.fill = GREY
                    c.font = Font(bold=True)
                current_row += 1

                # Cross-tab rows
                for gt_lbl in labels:
                    ws4.cell(row=current_row, column=1, value=gt_lbl).font = Font(bold=True)
                    for j, ai_lbl in enumerate(labels, 2):
                        count = sum(1 for t, p in zip(y_true, y_pred) if t == gt_lbl and p == ai_lbl)
                        cell = ws4.cell(row=current_row, column=j, value=count)
                        if gt_lbl == ai_lbl and count > 0:
                            cell.fill = GREEN
                        elif count > 0:
                            cell.fill = RED
                    current_row += 1

                current_row += 1  # blank row between fields

    print(f"  → Report written: {out_path}")
    return out_path
