#!/usr/bin/env python3
"""
Generalized F1/precision/recall scorer for one-row-per-study forms
(study characteristics, patient population, etc.).

Point it at any AI extraction CSV and any aligned GT Excel sheet — it
auto-infers a comparison strategy per shared column (numeric_exact /
exact_normalize / llm_judge), matches studies by fuzzy name, scores every
field, prints F1/precision/recall, and writes a color-coded comparison
workbook in the same "All Fields / Has Mismatch / <field>" layout as
outputs/full_studies/<corpus>/<model>/comparison_sheets/*.xlsx.

This wraps the same engine (eval/engine/{core,forms,config}) the per-review
scripts in eval/engine/forms/*.py use — it just skips writing a bespoke
FieldSpec list by hand when you only need a quick score.

NOTE: level-1 forms only (one row per study). Per-arm / per-outcome forms
(interventions, outcomes) need match_subrecords and belong in a proper
forms/<review>.py config — see eval/engine/forms/ibuprofen.py for the pattern.

Examples
--------
  python3 eval/score_form.py \\
      --ai "eval/sheets/ai sheets/full_studies/ibuprofen/claude/form_CD015432_—_Study_Characteristics_v3_long.csv" \\
      --gt "eval/sheets/gt sheets/ibuprofen.xlsx" --gt-sheet study_char

  python3 eval/score_form.py --ai my_extraction.csv --gt my_gt.xlsx --gt-sheet study_char \\
      --override age_value=numeric_tolerance:0.5 --out-dir /tmp/scores
"""

import argparse
import os
import re
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ENGINE_DIR = os.path.join(SCRIPT_DIR, "engine")
sys.path.insert(0, ENGINE_DIR)

import pandas as pd
from dotenv import load_dotenv

load_dotenv(os.path.join(SCRIPT_DIR, ".env"))

from config.field_config import (
    FieldSpec,
    STRATEGY_EXACT_NORMALIZE, STRATEGY_NUMERIC_EXACT, STRATEGY_NUMERIC_TOLERANCE,
    STRATEGY_LLM_JUDGE, STRATEGY_SKIP,
)
from config.nr_synonyms import is_nr
from core.cleaner import to_float
from core.loader import load_ai, load_gt_aligned
from forms.base_form import run_form

IDENTIFIER_NAMES = {"paper", "study_id", "refid", "authors_last_name", "id", "_gt_author"}
DIRECTIVE_TOKENS = {"extract", "optional", "neglect", "meta"}
VALID_STRATEGIES = {
    STRATEGY_EXACT_NORMALIZE, STRATEGY_NUMERIC_EXACT, STRATEGY_NUMERIC_TOLERANCE,
    STRATEGY_LLM_JUDGE, STRATEGY_SKIP,
}


def _looks_like_directive_row(gt_path: str, sheet: str) -> bool:
    """True if the first data row is a directive row (EXTRACT/OPTIONAL/NEGLECT/META),
    like the antibiotic/ibuprofen/periodontitis GT sheets use."""
    raw = pd.read_excel(gt_path, sheet_name=sheet, header=0, nrows=1)
    if raw.empty:
        return False
    vals = [str(v).strip().lower() for v in raw.iloc[0].tolist() if pd.notna(v)]
    if not vals:
        return False
    hits = sum(1 for v in vals if v in DIRECTIVE_TOKENS)
    return hits / len(vals) >= 0.5


def _infer_strategy(values, sample_n: int = 50) -> str:
    """Heuristic: mostly-numeric -> numeric_exact; low-cardinality short strings
    -> exact_normalize (categorical); otherwise llm_judge (free text)."""
    vals = [v for v in values if not is_nr(v)]
    if not vals:
        return STRATEGY_LLM_JUDGE
    sample = vals[:sample_n]
    numeric_hits = sum(1 for v in sample if to_float(v) is not None)
    if numeric_hits / len(sample) >= 0.8:
        return STRATEGY_NUMERIC_EXACT
    uniq = {str(v).strip().lower() for v in vals}
    avg_len = sum(len(str(v)) for v in sample) / len(sample)
    if len(uniq) <= max(6, round(0.25 * len(vals))) and avg_len <= 40:
        return STRATEGY_EXACT_NORMALIZE
    return STRATEGY_LLM_JUDGE


def _parse_overrides(raw_list):
    overrides = {}
    for item in raw_list or []:
        field_part, sep, strat_part = item.partition("=")
        if not sep:
            sys.exit(f"--override must be field=strategy[:tolerance], got: {item!r}")
        strat, _, tol = strat_part.partition(":")
        strat = strat.strip()
        if strat not in VALID_STRATEGIES:
            sys.exit(f"Unknown strategy '{strat}' in --override {item!r}. "
                      f"Valid: {sorted(VALID_STRATEGIES)}")
        overrides[field_part.strip()] = (strat, float(tol) if tol else 0.0)
    return overrides


def _default_out_dir(ai_path: str) -> str:
    """Mirror this repo's sheets/ai sheets/... -> outputs/... convention when
    the AI path lives under it; otherwise fall back to a sibling folder."""
    mapped = ai_path.replace("sheets/ai sheets", "outputs")
    base = mapped if mapped != ai_path else ai_path
    return os.path.join(os.path.dirname(base), "comparison_sheets")


def _write_comparison_workbook(comparison_df, fields, out_dir: str, form_name: str) -> str:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    out_path = os.path.join(out_dir, f"{form_name}_comparison.xlsx")

    F_STUDY  = PatternFill("solid", fgColor="F8F9FA")
    F_AI_HDR = PatternFill("solid", fgColor="E8A84A")
    F_GT_HDR = PatternFill("solid", fgColor="5EBB8A")
    F_MATCH  = PatternFill("solid", fgColor="B7EDD0")
    F_MISS   = PatternFill("solid", fgColor="F7B8B8")
    F_NA     = PatternFill("solid", fgColor="E8E8E8")
    HDR_FT   = Font(bold=True, color="FFFFFF")
    WRAP     = Alignment(wrap_text=True, vertical="top")
    _THIN    = Side(border_style="thin",   color="BDBDBD")
    _THICK   = Side(border_style="medium", color="6C6C6C")
    B_AI     = Border(left=_THICK, right=_THIN,  top=_THIN, bottom=_THIN)
    B_GT     = Border(left=_THIN,  right=_THICK, top=_THIN, bottom=_THIN)
    B_STUDY  = Border(left=_THICK, right=_THICK, top=_THIN, bottom=_THIN)

    scored = [f for f in fields if f"AI_{f.ai_col}" in comparison_df.columns]
    strat_by_field = {f.ai_col: f.strategy for f in scored}
    scored_cols = [f.ai_col for f in scored]

    def _write_sheet(wb, name, df, cols):
        ws = wb.create_sheet(name[:31])
        h1 = ws.cell(1, 1, "Study")
        h1.fill = F_STUDY; h1.font = Font(bold=True); h1.border = B_STUDY
        col = 2
        for f_col in cols:
            a = ws.cell(1, col, f_col); a.fill = F_AI_HDR; a.font = HDR_FT; a.border = B_AI
            g = ws.cell(1, col + 1);    g.fill = F_GT_HDR;                  g.border = B_GT
            col += 2
        h2 = ws.cell(2, 1, "Strategy →")
        h2.fill = F_STUDY; h2.font = Font(italic=True); h2.border = B_STUDY
        col = 2
        for f_col in cols:
            s = strat_by_field.get(f_col, "")
            a = ws.cell(2, col,     f"AI  ({s})"); a.fill = F_AI_HDR; a.font = HDR_FT; a.border = B_AI
            g = ws.cell(2, col + 1, f"GT  ({s})"); g.fill = F_GT_HDR; g.font = HDR_FT; g.border = B_GT
            col += 2
        for r_idx, (_, row) in enumerate(df.iterrows(), start=3):
            study = row.get("AI_study") or row.get("GT_study") or ""
            sc = ws.cell(r_idx, 1, study); sc.fill = F_STUDY; sc.border = B_STUDY
            col = 2
            for f_col in cols:
                ai_v = row.get(f"AI_{f_col}", "")
                gt_v = row.get(f"GT_{f_col}", "")
                m    = row.get(f"MATCH_{f_col}", "")
                fill = F_MATCH if m == "YES" else F_MISS if m == "NO" else F_NA
                a = ws.cell(r_idx, col,     ai_v); a.fill = fill; a.alignment = WRAP; a.border = B_AI
                g = ws.cell(r_idx, col + 1, gt_v); g.fill = fill; g.alignment = WRAP; g.border = B_GT
                col += 2
        ws.column_dimensions["A"].width = 26
        for c in range(2, 2 + 2 * len(cols)):
            ws.column_dimensions[get_column_letter(c)].width = 30
        ws.freeze_panes = "B3"
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 22

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_sheet(wb, "All Fields", comparison_df, scored_cols)

    match_cols = [f"MATCH_{c}" for c in scored_cols if f"MATCH_{c}" in comparison_df.columns]
    mismatch_mask = (comparison_df[match_cols].eq("NO").any(axis=1)
                      if match_cols else pd.Series([False] * len(comparison_df)))
    _write_sheet(wb, "Has Mismatch", comparison_df[mismatch_mask], scored_cols)

    for f_col in scored_cols:
        _write_sheet(wb, f_col, comparison_df, [f_col])

    wb.save(out_path)
    return out_path


def main():
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--ai", required=True, help="AI extraction CSV (one row per study).")
    parser.add_argument("--gt", required=True, help="Aligned GT Excel workbook.")
    parser.add_argument("--gt-sheet", default="study_char", help="Sheet name in the GT workbook.")
    parser.add_argument("--gt-key-col", default=None,
                         help="Study-id column in the GT sheet (default: auto — study_id, else Paper).")
    parser.add_argument("--gt-skiprows", default=None,
                         help="Comma-separated row indices to skip after the header "
                              "(default: auto-detected EXTRACT/OPTIONAL/NEGLECT/META directive row).")
    parser.add_argument("--form-name", default=None,
                         help="Name for this form/run (default: derived from GT filename + sheet).")
    parser.add_argument("--out-dir", default=None,
                         help="Where to write the comparison workbook "
                              "(default: mirrors this repo's sheets/ai sheets -> outputs convention, "
                              "else <ai dir>/comparison_sheets).")
    parser.add_argument("--no-llm", action="store_true",
                         help="Disable the LLM judge (llm_judge fields are excluded, not scored).")
    parser.add_argument("--override", action="append", default=[],
                         help="Override an inferred strategy, e.g. "
                              "--override age_value=numeric_tolerance:0.5 (repeatable).")
    args = parser.parse_args()

    ai_path = os.path.abspath(args.ai)
    gt_path = os.path.abspath(args.gt)
    if not os.path.exists(ai_path):
        sys.exit(f"AI file not found: {ai_path}")
    if not os.path.exists(gt_path):
        sys.exit(f"GT file not found: {gt_path}")

    sheet_names = pd.ExcelFile(gt_path).sheet_names
    if args.gt_sheet not in sheet_names:
        sys.exit(f"Sheet '{args.gt_sheet}' not in {gt_path}. Available: {sheet_names}")

    gt_key_col = args.gt_key_col
    if gt_key_col is None:
        header_cols = pd.read_excel(gt_path, sheet_name=args.gt_sheet, nrows=0).columns
        gt_key_col = "study_id" if "study_id" in header_cols else "Paper"
        print(f"[auto] gt-key-col = {gt_key_col!r}")

    if args.gt_skiprows is not None:
        gt_skiprows = [int(x) for x in args.gt_skiprows.split(",") if x.strip() != ""]
    else:
        gt_skiprows = [1] if _looks_like_directive_row(gt_path, args.gt_sheet) else None
        print(f"[auto] gt-skiprows = {gt_skiprows}")

    form_name = args.form_name or re.sub(
        r"[^a-z0-9]+", "_",
        f"{os.path.splitext(os.path.basename(gt_path))[0]}_{args.gt_sheet}".lower()
    ).strip("_")

    ai_df = load_ai(ai_path, "csv")
    gt_df = load_gt_aligned(gt_path, args.gt_sheet, key_col=gt_key_col, skiprows=gt_skiprows)

    overrides = _parse_overrides(args.override)

    fields = []
    print(f"\nInferred field strategies for form '{form_name}':")
    for col in gt_df.columns:
        if col.strip().lower() in IDENTIFIER_NAMES:
            continue
        if col not in ai_df.columns:
            continue
        if col in overrides:
            strat, tol = overrides[col]
        else:
            strat, tol = _infer_strategy(gt_df[col].tolist()), 0.0
        fields.append(FieldSpec(ai_col=col, strategy=strat, tolerance=tol))
        print(f"  {col:<32} -> {strat}" + (f" (tol={tol})" if tol else ""))

    if not fields:
        sys.exit("No shared, scoreable columns found between the AI sheet and the GT sheet.")

    out_dir = args.out_dir or _default_out_dir(ai_path)
    os.makedirs(out_dir, exist_ok=True)
    metrics_dir = os.path.dirname(out_dir) if os.path.basename(out_dir) == "comparison_sheets" else out_dir
    os.environ["EVAL_METRICS_DIR"] = metrics_dir

    result = run_form(
        form_name=form_name,
        fields=fields,
        ai_path=ai_path,
        ai_format="csv",
        gt_section=args.gt_sheet,
        level2=False,
        use_llm=not args.no_llm,
        gt_aligned_sheet=args.gt_sheet,
        aligned_gt_path=gt_path,
        gt_key_col=gt_key_col,
        gt_skiprows=gt_skiprows,
        canonical_form=form_name,
    )

    print("\n=== Field metrics ===")
    print(result["stats_df"].to_string(index=False))
    print("\n=== Form summary ===")
    for k, v in result["summary"].items():
        print(f"  {k}: {v}")

    out_path = _write_comparison_workbook(result["comparison_df"], fields, out_dir, form_name)
    print(f"\n✓ Comparison workbook: {out_path}")
    print(f"✓ Metrics workbook:     {os.path.join(metrics_dir, 'all_forms_metrics.xlsx')}")


if __name__ == "__main__":
    main()
