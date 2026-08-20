"""Read the text out of a review-planning document — protocol, PROSPERO record,
eligibility-criteria table — for the review-scope suggester.

Deliberately NOT the document pipeline. ``pdf_processing_service`` runs
marker/Datalab under Celery to produce per-page markdown with bounding boxes,
because extraction needs source linking; that is minutes of GPU work and a row
in ``documents``. A protocol is read once, never stored, and only needs a flat
string — so this is a synchronous PyMuPDF / python-docx read with a hard
character budget.

Nothing here truncates silently: ``assemble_document`` returns a per-file report
saying exactly how much of each file was read, and the endpoint passes it to the
browser.
"""

from __future__ import annotations

import io
import logging
import os
import re
from typing import Dict, List, Sequence, Tuple

logger = logging.getLogger(__name__)

# Enough for a 40-page protocol. A published review PDF beyond this is truncated
# and *said* to be truncated.
MAX_TOTAL_CHARS = 200_000

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}

# Rejected with a sentence the user can act on rather than an obscure parse
# failure: there is no libreoffice on this box, so legacy binary .doc cannot be
# converted.
_LEGACY_WORD = {".doc"}

_NON_WORD = re.compile(r"[^a-z0-9 ]+")
_WS = re.compile(r"\s+")


class UnsupportedDocument(ValueError):
    """Raised with a message meant for the user, not the log."""


def extension_of(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def normalize_for_match(text: str) -> str:
    """Lowercase, drop punctuation, collapse whitespace.

    PDF extraction breaks words across lines, turns hyphens into newlines and
    pads table cells with runs of spaces, so a quote the model copied correctly
    will not match the source byte-for-byte. Normalising both sides this hard is
    what makes the evidence check in ``scope_suggestion`` usable at all.
    """
    return _WS.sub(" ", _NON_WORD.sub(" ", (text or "").lower())).strip()


def extract_text(filename: str, data: bytes) -> str:
    """Flat text for one uploaded file. Raises UnsupportedDocument by design."""
    ext = extension_of(filename)
    if ext in _LEGACY_WORD:
        raise UnsupportedDocument(
            f"{filename}: legacy .doc files cannot be read here — "
            f"save it as .docx or PDF and try again."
        )
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocument(
            f"{filename}: unsupported file type. Upload a PDF, .docx, .md or .txt."
        )
    if ext == ".pdf":
        return _from_pdf(data)
    if ext == ".docx":
        return _from_docx(data)
    return _from_plain(data)


def _from_pdf(data: bytes) -> str:
    import fitz  # PyMuPDF — same in-memory open as services/pdf_cleaner.py

    pages: List[str] = []
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        for page in doc:
            pages.append(page.get_text("text") or "")
    finally:
        doc.close()
    return "\n\n".join(pages)


def _from_docx(data: bytes) -> str:
    """Paragraphs AND table cells.

    Eligibility criteria in a protocol are nearly always a two-column table
    (Include | Exclude). python-docx's ``.paragraphs`` skips table content
    entirely, so reading only paragraphs would drop the single most important
    section in the file.
    """
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise UnsupportedDocument(
            "Word documents cannot be read on this server — upload a PDF instead."
        ) from exc

    document = docx.Document(io.BytesIO(data))
    parts: List[str] = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            # A merged cell reports the same text in every column it spans;
            # collapsing keeps "Inclusion criteria" from arriving four times.
            deduped = list(dict.fromkeys([c for c in cells if c]))
            if deduped:
                parts.append(" | ".join(deduped))
    return "\n".join(parts)


def _from_plain(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def assemble_document(
    files: Sequence[Tuple[str, bytes]],
    max_total_chars: int = MAX_TOTAL_CHARS,
) -> Tuple[str, List[Dict[str, object]]]:
    """Concatenate several files into one prompt payload, budget-capped.

    Returns ``(document_text, per_file_report)``. The budget is spent in the
    order given, so a truncation always falls at the END of the payload and the
    report names exactly which file lost what — the reviewer can then re-upload
    that file on its own rather than wondering why an outcome went missing.
    """
    chunks: List[str] = []
    report: List[Dict[str, object]] = []
    remaining = max_total_chars

    for filename, data in files:
        text = extract_text(filename, data).strip()
        total = len(text)
        kept = "" if remaining <= 0 else text[:remaining]
        remaining -= len(kept)
        report.append(
            {
                "filename": filename,
                "chars_total": total,
                "chars_read": len(kept),
                "truncated": len(kept) < total,
                "empty": total == 0,
            }
        )
        if kept:
            chunks.append(f"--- {filename} ---\n{kept}")

    return "\n\n".join(chunks), report
